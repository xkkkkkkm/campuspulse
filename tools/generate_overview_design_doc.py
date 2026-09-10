from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "101-概要设计说明书(模版).docx"
OUTPUT = ROOT / "校园活动智能推荐与组队平台-概要设计说明书.docx"
SCHEMA = ROOT / "backend/src/main/resources/schema.sql"
PLANTUML_JAR = ROOT / "tools/plantuml.jar"
ASSET_DIR = ROOT / "docs/overview_design/assets"
PUML_DIR = ROOT / "docs/overview_design/puml"


@dataclass
class Column:
    name: str
    data_type: str
    nullable: str
    default: str
    constraints: list[str]


@dataclass
class TableDef:
    name: str
    columns: list[Column]
    table_constraints: list[str]


GENERIC_FIELD_DESC = {
    "id": "主键标识，使用自增长整型唯一标识记录。",
    "user_id": "关联用户主键，用于绑定业务记录所属用户。",
    "activity_id": "关联活动主键，用于建立与活动实体的关联关系。",
    "team_id": "关联队伍主键，用于建立与组队实体的关联关系。",
    "sender_id": "消息发送方用户标识。",
    "receiver_id": "消息接收方用户标识。",
    "peer_id": "私聊对端用户标识。",
    "tag_id": "关联标签主键。",
    "creator_id": "创建人用户标识。",
    "organizer_id": "活动发布者或组织者用户标识。",
    "title": "标题信息，用于前端卡片、详情页和消息提示展示。",
    "name": "名称字段，用于描述标签或分类名称。",
    "type": "类型枚举字段，用于区分业务类别。",
    "status": "状态枚举字段，用于控制业务流程及页面展示。",
    "role": "角色枚举字段，用于表示用户、成员或权限身份。",
    "content": "正文内容字段，用于保存消息、通知或说明文本。",
    "description": "详细描述字段，用于展示业务对象的详细说明。",
    "message": "补充说明信息，多用于申请附言或接口消息。",
    "location": "地点信息，用于描述活动举办地点。",
    "cover_url": "封面图片访问路径。",
    "avatar_url": "头像图片访问路径。",
    "max_participants": "活动人数上限，0 表示不限制。",
    "max_members": "队伍人数上限，0 表示不限制。",
    "start_time": "业务开始时间。",
    "end_time": "业务结束时间。",
    "created_at": "记录创建时间。",
    "updated_at": "记录更新时间。",
    "joined_at": "成员加入队伍时间。",
    "event_time": "行为事件发生时间。",
    "read_flag": "通知是否已读标识，0 表示未读，1 表示已读。",
    "code_hash": "验证码摘要值，不以明文形式存储。",
    "expires_at": "验证码失效时间。",
    "consumed_at": "验证码被使用时间。",
    "attempt_count": "验证码校验失败次数。",
    "username": "用户登录账号。",
    "student_no": "学号，用于校园身份识别。",
    "password_hash": "密码摘要，采用单向哈希算法保存。",
    "nickname": "用户昵称。",
    "college": "学院信息。",
    "campus": "校区信息。",
    "major": "专业信息。",
    "grade": "年级信息。",
    "education_level": "培养层次或学历层次。",
    "bio": "个人简介。",
    "email": "邮箱地址。",
    "email_verified": "邮箱是否已完成验证。",
    "phone": "联系电话。",
    "chat_enabled": "活动群聊是否开启。",
    "teaming_enabled": "活动是否允许关联组队。",
    "version": "乐观锁版本号或并发控制字段。",
    "real_name": "报名时填写的真实姓名。",
    "intro": "报名或申请时填写的补充自我介绍。",
    "last_read_id": "用户在某个私聊会话中的最后已读消息 ID。",
    "event_type": "行为事件类型，如点击、收藏、报名等。",
    "extra_json": "行为附加上下文，使用 JSON 字符串保存。",
    "weight": "运营推荐权重，数值越大排序越靠前。",
    "purpose": "验证码用途标识，如注册、登录或重置密码。",
}


FIELD_DESC_OVERRIDE = {
    "tags.type": "标签分类，当前系统主要用于首页分类、兴趣标签和扩展标签管理。",
    "users.role": "用户角色，系统中主要包括普通用户、活动组织者和管理员。",
    "users.status": "用户状态，1 表示启用，0 表示禁用。",
    "activities.status": "活动发布状态，主要用于控制上架、下架和删除逻辑。",
    "activities.audit_status": "活动审核状态，控制活动是否能够对外展示。",
    "registrations.status": "报名状态，主要包括 APPLIED、APPROVED、REJECTED、CANCELLED 等。",
    "team_member.role": "队伍成员角色，如 CREATOR、MEMBER。",
    "team_member.status": "成员状态，ACTIVE 表示有效成员。",
    "team_join_request.status": "入队申请状态，主要包括 PENDING、APPROVED、REJECTED。",
    "notifications.type": "通知业务类型，用于驱动消息分类和前端展示。",
    "activities.version": "预留的乐观锁或版本控制字段，为后续并发控制扩展提供支持。",
    "teams.version": "预留的乐观锁或版本控制字段，为队伍并发修改预留。",
}


TABLE_CN = {
    "users": "用户表",
    "tags": "标签表",
    "sms_code": "短信验证码表",
    "email_code": "邮箱验证码表",
    "user_interest": "用户兴趣关联表",
    "activities": "活动表",
    "activity_tag": "活动标签关联表",
    "registrations": "活动报名表",
    "favorites": "活动收藏表",
    "teams": "队伍表",
    "team_member": "队伍成员表",
    "team_join_request": "入队申请表",
    "messages": "队伍聊天消息表",
    "dm_message": "私聊消息表",
    "dm_read_state": "私聊已读状态表",
    "activity_chat_message": "活动群聊消息表",
    "notifications": "通知表",
    "event_log": "行为事件日志表",
    "ops_featured_activity": "运营推荐活动表",
}


TABLE_PURPOSE = {
    "users": "保存平台账号、身份、联系方式和角色权限等基础信息，是全系统的核心主数据实体。",
    "tags": "保存活动分类、兴趣标签与运营标签，是推荐、筛选和画像构建的重要基础数据。",
    "sms_code": "保存短信验证码签发、有效期和尝试次数，用于移动端或扩展短信校验场景。",
    "email_code": "保存邮箱验证码签发、有效期和尝试次数，用于邮箱登录、绑定、修改密码等场景。",
    "user_interest": "维护用户与兴趣标签的多对多关系，为个性化推荐和组队匹配提供输入。",
    "activities": "保存活动主数据，包括基本信息、时间地点、审核状态、人数上限和互动开关。",
    "activity_tag": "维护活动与标签之间的多对多关系，用于活动筛选、推荐排序和标签展示。",
    "registrations": "保存用户活动报名记录以及审批状态，用于报名管理与人数统计。",
    "favorites": "保存用户收藏活动记录，用于个人中心展示与推荐权重计算。",
    "teams": "保存与活动关联或独立存在的队伍主数据，用于组队大厅和队伍详情展示。",
    "team_member": "保存队伍成员及其角色信息，是队伍权限控制和聊天访问控制的依据。",
    "team_join_request": "保存用户申请加入队伍的申请记录与审核状态，用于队长审批流程。",
    "messages": "保存队伍频道聊天消息，支持团队协作沟通与实时推送。",
    "dm_message": "保存两名用户之间的私聊消息，是站内即时沟通能力的核心数据表。",
    "dm_read_state": "记录用户在不同私聊线程中的已读游标，用于未读数与消息已读管理。",
    "activity_chat_message": "保存活动群聊消息，支撑活动预热、交流和提醒类互动。",
    "notifications": "保存系统通知、审核结果、报名结果、入队申请等消息提醒记录。",
    "event_log": "保存用户点击、收藏、报名等行为事件，是推荐排序和运营分析的基础数据。",
    "ops_featured_activity": "保存运营推荐位配置，用于首页精选展示和推荐流权重增强。",
}


DIAGRAMS = {
    "overall-architecture.puml": r"""@startuml
skinparam defaultFontName "PingFang SC"
skinparam shadowing false
skinparam packageStyle rectangle
skinparam componentStyle rectangle
skinparam rectangle {
  RoundCorner 15
  BackgroundColor #F7F9FC
  BorderColor #3B4A6B
}
title 图3-1 系统总体架构图

rectangle "Web 用户平台" as web_user {
  component "首页/活动大厅" as user_home
  component "组队大厅/消息中心" as user_team
  component "个人中心" as user_profile
}

rectangle "Web 管理后台" as admin_web {
  component "用户管理" as admin_user
  component "活动审核" as admin_activity
  component "推荐位/标签运营" as admin_ops
}

rectangle "CampusPulse 后端服务\nSpring Boot 3.1.4" as backend {
  component "认证与权限控制"
  component "活动与报名管理"
  component "推荐服务"
  component "组队协作服务"
  component "消息与通知服务"
  component "上传与支撑服务"
}

database "MySQL 8.x\n业务数据库" as mysql
rectangle "文件存储\n封面/头像" as storage
cloud "SMTP 邮件服务" as smtp
cloud "AI 外部服务\n标签提取/推荐辅助/组队建议" as ai
cloud "Dify 知识库 Chatflow\n智能客服问答" as dify

web_user --> backend : REST/JSON\nSSE
admin_web --> backend : REST/JSON
backend --> mysql : JDBC
backend --> storage : 文件读写
backend --> smtp : 邮件验证码
backend --> ai : HTTP/JSON
backend --> dify : HTTP/JSON
@enduml
""",
    "functional-architecture.puml": r"""@startuml
skinparam defaultFontName "PingFang SC"
skinparam shadowing false
skinparam packageStyle rectangle
title 图3-2 功能架构图

rectangle "校园活动智能推荐与组队平台" {
  package "1 用户与认证模块" {
    [账号注册登录]
    [兴趣标签维护]
    [资料维护]
  }
  package "2 活动发现与活动管理模块" {
    [活动浏览检索]
    [活动详情与收藏]
    [报名与审批]
    [活动发布维护]
  }
  package "3 智能推荐模块" {
    [首页推荐流]
    [兴趣画像计算]
    [运营加权排序]
  }
  package "4 组队协作模块" {
    [队伍创建维护]
    [入队申请审批]
    [队员展示]
  }
  package "5 消息通信模块" {
    [队伍聊天]
    [活动群聊]
    [用户私聊]
    [SSE 实时推送]
  }
  package "6 通知与个人中心模块" {
    [通知中心]
    [我的报名]
    [我的收藏]
    [我的队伍]
    [智能客服]
  }
  package "7 管理后台模块" {
    [用户管理]
    [活动审核]
    [标签管理]
    [推荐位管理]
    [统计看板]
  }
  package "8 平台支撑模块" {
    [统一响应与异常处理]
    [文件上传]
    [行为埋点]
    [定时清理任务]
  }
}

[账号注册登录] ..> [统一响应与异常处理]
[活动浏览检索] ..> [首页推荐流]
[报名与审批] ..> [通知中心]
[入队申请审批] ..> [通知中心]
[队伍聊天] ..> [SSE 实时推送]
[活动群聊] ..> [SSE 实时推送]
[用户私聊] ..> [SSE 实时推送]
[首页推荐流] ..> [兴趣画像计算]
[首页推荐流] ..> [运营加权排序]
[智能客服] ..> [统一响应与异常处理]
@enduml
""",
    "data-architecture.puml": r"""@startuml
skinparam defaultFontName "PingFang SC"
skinparam shadowing false
hide circle
skinparam linetype ortho
title 图3-3 核心数据实体关系图

entity "users" as users
entity "tags" as tags
entity "user_interest" as user_interest
entity "activities" as activities
entity "activity_tag" as activity_tag
entity "registrations" as registrations
entity "favorites" as favorites
entity "teams" as teams
entity "team_member" as team_member
entity "team_join_request" as team_join_request
entity "messages" as messages
entity "dm_message" as dm_message
entity "dm_read_state" as dm_read_state
entity "activity_chat_message" as activity_chat_message
entity "notifications" as notifications
entity "event_log" as event_log
entity "ops_featured_activity" as ops_featured_activity
entity "email_code" as email_code
entity "sms_code" as sms_code

users ||--o{ user_interest
tags ||--o{ user_interest
users ||--o{ activities : 发布
activities ||--o{ activity_tag
tags ||--o{ activity_tag
users ||--o{ registrations
activities ||--o{ registrations
users ||--o{ favorites
activities ||--o{ favorites
activities ||--o{ teams
users ||--o{ teams : 创建
teams ||--o{ team_member
users ||--o{ team_member
teams ||--o{ team_join_request
users ||--o{ team_join_request
teams ||--o{ messages
users ||--o{ messages
users ||--o{ dm_message : sender
users ||--o{ dm_message : receiver
users ||--o{ dm_read_state
activities ||--o{ activity_chat_message
users ||--o{ activity_chat_message
users ||--o{ notifications
users ||--o{ event_log
activities ||--o{ event_log
activities ||--|| ops_featured_activity
email_code }o--|| users : 逻辑关联
sms_code }o--|| users : 逻辑关联
@enduml
""",
    "deployment-architecture.puml": r"""@startuml
skinparam defaultFontName "PingFang SC"
skinparam shadowing false
skinparam node {
  RoundCorner 15
  BackgroundColor #F6FAFF
  BorderColor #2F5C8A
}
title 图3-4 部署架构图

node "用户终端" as client {
  artifact "Chrome / Edge / Safari"
}

node "Web 前端部署" as frontend {
  artifact "HTML/CSS/JavaScript 静态资源"
  artifact "原型页面与上传资源目录"
}

node "应用服务主机" as app {
  artifact "CampusPulse API\nSpring Boot Jar"
  artifact "SSE 实时推送服务"
  artifact "定时清理任务"
}

database "MySQL 8.x" as db
node "文件存储目录" as fs {
  artifact "avatars/"
  artifact "covers/"
}
cloud "SMTP 服务" as smtp
cloud "AI 服务" as ai
cloud "Dify Chatflow 服务" as dify

client --> frontend : HTTPS
client --> app : REST/JSON\nSSE
frontend --> app : API 调用
app --> db : JDBC
app --> fs : 本地文件 I/O
app --> smtp : SMTP/TLS
app --> ai : HTTP/JSON
app --> dify : HTTP/JSON
@enduml
""",
}


INTERNAL_INTERFACES = [
    ("认证鉴权接口", "用户与认证模块", [
        ("POST", "/api/auth/register", "普通账号注册，写入 users 并签发 Token。"),
        ("POST", "/api/auth/login", "账号密码登录，返回 ApiResponse<AuthResponse>。"),
        ("POST", "/api/auth/login-email", "邮箱验证码登录，降低密码遗忘带来的使用门槛。"),
        ("GET", "/api/auth/me", "读取当前登录用户资料和角色信息。"),
        ("PUT", "/api/auth/me/interests", "维护用户兴趣标签，驱动推荐画像。"),
    ]),
    ("活动与报名接口", "活动发现与活动管理模块", [
        ("GET", "/api/activities", "按关键字、标签、分页查询活动列表。"),
        ("POST", "/api/activities", "发布活动并同步活动标签。"),
        ("GET", "/api/activities/{id}", "查看活动详情、标签、收藏状态和聊天开关。"),
        ("POST", "/api/activities/{id}/register", "提交活动报名申请。"),
        ("POST", "/api/activities/{activityId}/registrations/{registrationId}/approve", "组织者审批通过报名。"),
    ]),
    ("推荐计算接口", "智能推荐模块", [
        ("GET", "/api/recommendations/feed", "基于兴趣标签、历史参加活动、历史组队参与、行为日志、时间衰减与运营权重生成推荐流。"),
        ("POST", "/api/events", "记录 CLICK、FAVORITE、REGISTER 等行为，为推荐算法提供输入，并补充用户历史兴趣画像。"),
    ]),
    ("组队协作接口", "组队协作模块", [
        ("GET", "/api/teams", "按活动或全局查询开放中的队伍列表。"),
        ("POST", "/api/teams", "创建队伍并自动写入队长成员关系。"),
        ("POST", "/api/teams/{id}/join", "提交入队申请。"),
        ("POST", "/api/teams/{teamId}/requests/{requestId}/approve", "队长审批通过入队请求。"),
        ("GET", "/api/activities/{id}/teams", "在活动详情下展示关联队伍。"),
    ]),
    ("消息通信接口", "消息通信模块", [
        ("GET", "/api/teams/{teamId}/messages", "获取队伍频道历史消息。"),
        ("POST", "/api/teams/{teamId}/messages", "发送队伍消息并触发实时通知。"),
        ("GET", "/api/dm/{peerId}/messages", "获取用户私聊记录。"),
        ("POST", "/api/dm/{peerId}/messages", "发送私聊消息。"),
        ("GET", "/api/activities/{id}/chat/messages", "获取活动群聊消息。"),
        ("POST", "/api/activities/{id}/chat/messages", "发送活动群聊消息。"),
        ("GET", "/api/realtime/chat/stream", "建立 SSE 长连接，向队伍、私聊、活动频道推送消息。"),
    ]),
    ("通知与个人中心接口", "通知与个人中心模块", [
        ("GET", "/api/notifications", "按已读状态分页查询通知中心。"),
        ("POST", "/api/notifications/read-all", "一键标记全部通知为已读。"),
        ("GET", "/api/profile/registrations", "查看我的报名记录。"),
        ("GET", "/api/profile/favorites", "查看我的收藏活动。"),
        ("GET", "/api/profile/teams", "查看我的队伍与聊天入口。"),
    ]),
    ("智能客服接口", "通知与个人中心模块", [
        ("HTTP", "Dify Chatflow API", "通过知识库 Chatflow 解答用户使用问题，返回步骤说明、功能入口提示和常见问题答案。"),
    ]),
    ("后台运营接口", "管理后台模块", [
        ("GET", "/api/admin/stats", "汇总用户、活动、消息等统计数据。"),
        ("GET", "/api/admin/users", "分页检索用户列表。"),
        ("POST", "/api/admin/users/{id}/status", "启停用用户账号。"),
        ("POST", "/api/admin/activities/{id}/audit", "审核活动内容并通知发布者。"),
        ("POST", "/api/admin/featured", "维护精选活动及推荐权重。"),
    ]),
    ("文件上传接口", "平台支撑模块", [
        ("POST", "/api/upload/avatar", "上传头像文件，返回访问路径。"),
        ("POST", "/api/upload/cover", "上传活动封面文件，返回访问路径。"),
    ]),
]


EXTERNAL_INTERFACES = [
    ("Web 前端与后端 REST 接口", "HTTPS + REST/JSON", "浏览器前端发起 HTTP 请求，后端统一返回 ApiResponse<T> 结构。", "Token 鉴权、活动浏览、报名、组队、个人中心、后台管理。"),
    ("浏览器与后端 SSE 接口", "HTTP + Server-Sent Events", "前端在消息场景下建立单向实时订阅通道，服务端按频道推送实时消息事件。", "队伍聊天、活动群聊、私聊实时更新。"),
    ("后端与 MySQL 数据库接口", "JDBC", "Spring Boot 通过 JdbcTemplate 访问 MySQL，完成业务数据读写与事务处理。", "用户、活动、组队、消息、推荐与日志持久化。"),
    ("后端与 SMTP 邮件服务接口", "SMTP/TLS", "邮箱验证码服务通过 SMTP 发送验证码邮件，并以 email_code 表管理验证码生命周期。", "邮箱注册、邮箱登录、找回密码、邮箱变更。"),
    ("后端与 AI 服务接口", "HTTP/JSON", "以活动标题、活动描述、用户兴趣画像或活动上下文为输入，获取标签提取、推荐辅助和组队建议结果。", "活动标签自动提取、推荐增强、智能组队建议。"),
    ("后端与 Dify Chatflow 接口", "HTTP/JSON", "后端将用户问题、页面上下文和知识库检索条件发送至 Dify Chatflow，获取智能客服回答结果。", "用户使用问题解答、功能指引、常见问题咨询。"),
    ("后端与文件存储接口", "本地文件 I/O + URL 映射", "上传模块将头像和封面写入静态资源目录，并返回可访问路径给前端。", "头像上传、活动封面上传。"),
]


ERROR_ROWS = [
    ("参数错误", "请求参数缺失、格式错误或业务校验失败", "HTTP 400 + ApiResponse(false, message, data)", "前端提示用户修正输入，服务端不写入业务数据。"),
    ("未认证", "Token 缺失、失效或解析失败", "HTTP 401 + ApiResponse(false, '未登录或登录已过期', null)", "前端跳转登录页或触发重新登录。"),
    ("无权限", "角色不符或资源不归属当前用户", "HTTP 403 + ApiResponse(false, message, null)", "阻断操作并记录必要日志。"),
    ("资源不存在", "活动、队伍、用户或消息查询为空", "HTTP 404 + ApiResponse(false, message, null)", "前端展示空态或提示资源已失效。"),
    ("状态冲突", "重复报名、重复申请入队、审核状态不允许操作", "HTTP 409 + ApiResponse(false, message, null)", "提示当前对象状态已变化，要求用户刷新后重试。"),
    ("外部服务失败", "AI 服务或邮件服务不可用", "HTTP 200/500 + 业务失败消息", "执行降级策略，保底使用本地规则、默认标签或稍后重试机制。"),
    ("系统异常", "未预料异常、数据库故障或 I/O 故障", "HTTP 500 + ApiResponse(false, '系统繁忙，请稍后再试', data)", "记录日志、限制异常扩散，并通过运维手段恢复。"),
]


MODULE_MATRIX_ROWS = [
    ("用户注册登录", "Y", "", "", "", "", "Y", "", "Y"),
    ("兴趣标签维护", "Y", "", "Y", "", "", "Y", "", ""),
    ("活动检索与详情", "", "Y", "Y", "", "", "", "", ""),
    ("活动收藏与报名", "", "Y", "Y", "", "Y", "Y", "", "Y"),
    ("活动发布与管理", "", "Y", "", "", "Y", "", "", "Y"),
    ("队伍创建与申请", "", "", "", "Y", "Y", "Y", "", "Y"),
    ("队伍/活动聊天", "", "", "", "Y", "Y", "", "", "Y"),
    ("通知中心", "", "", "", "", "", "Y", "", "Y"),
    ("智能客服问答", "", "", "", "", "", "Y", "", "Y"),
    ("后台审核与推荐位", "", "", "Y", "", "", "", "Y", "Y"),
    ("文件上传与行为埋点", "", "", "", "", "", "", "", "Y"),
]


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def ensure_heading_style(doc: Document, level: int) -> str:
    name = f"Heading {level}"
    try:
        doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.quick_style = True
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.bold = True
        style.font.size = Pt(16 if level == 1 else 14 if level == 2 else 12)
        ppr = style._element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level - 1))
        ppr.append(outline)
    return name


def remove_paragraph(paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)


def clear_after(doc: Document, start_index: int) -> None:
    for paragraph in list(doc.paragraphs[start_index:]):
        remove_paragraph(paragraph)


def set_run_font(run, size: float = 12, bold: bool = False, east_asia: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def apply_para_format(paragraph, heading_level: int | None = None) -> None:
    fmt = paragraph.paragraph_format
    if heading_level is None:
        fmt.first_line_indent = Cm(0.74)
        fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        fmt.space_after = Pt(0)
        fmt.space_before = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        fmt.first_line_indent = Cm(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        fmt.space_before = Pt(6)
        fmt.space_after = Pt(3)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_body_paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run1 = p.add_run(bold_prefix)
        set_run_font(run1, 12, True)
        run2 = p.add_run(text[len(bold_prefix):])
        set_run_font(run2, 12, False)
    else:
        run = p.add_run(text)
        set_run_font(run, 12, False)
    apply_para_format(p)


def add_heading(doc: Document, text: str, level: int) -> None:
    style = ensure_heading_style(doc, level)
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, 16, True, east_asia="黑体")
    elif level == 2:
        set_run_font(run, 14, True, east_asia="黑体")
    else:
        set_run_font(run, 12, True, east_asia="黑体")
    apply_para_format(p, level)


def add_inline_title(doc: Document, title: str, body: str) -> None:
    p = doc.add_paragraph()
    r1 = p.add_run(title)
    set_run_font(r1, 12, True)
    r2 = p.add_run(body)
    set_run_font(r2, 12, False)
    apply_para_format(p)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, 11, False)
    apply_para_format(p)
    p.paragraph_format.first_line_indent = Cm(0)


def style_cell(cell, bold: bool = False, center: bool = False, font_size: float = 10.5) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for paragraph in cell.paragraphs:
        if not paragraph.runs:
            run = paragraph.add_run("")
            set_run_font(run, font_size, bold)
        for run in paragraph.runs:
            set_run_font(run, font_size, bold)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "000000")


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths: list[float] | None = None,
    body_font_size: float = 10.5,
    header_font_size: float = 10.5,
):
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = str(text)
        style_cell(hdr[i], bold=True, center=True, font_size=header_font_size)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            style_cell(cells[i], center=False, font_size=body_font_size)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    return table


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "右键更新目录"
    fld_separate.append(text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_separate, fld_end])


def parse_schema(schema_text: str) -> list[TableDef]:
    tables: list[TableDef] = []
    blocks = re.findall(r"CREATE TABLE IF NOT EXISTS (\w+)\s*\((.*?)\)\s*ENGINE=", schema_text, re.S)
    for table_name, body in blocks:
        columns: list[Column] = []
        table_constraints: list[str] = []
        raw_lines = [line.strip().rstrip(",") for line in body.splitlines() if line.strip()]
        for raw in raw_lines:
            if raw.upper().startswith(("PRIMARY KEY", "UNIQUE KEY", "KEY ", "CONSTRAINT", "--")):
                table_constraints.append(raw)
                continue
            if raw.startswith(")") or raw.startswith("/*"):
                continue
            m = re.match(r"(\w+)\s+(.+)", raw)
            if not m:
                continue
            name, remainder = m.groups()
            constraint_tokens = []
            keywords = [" NOT NULL", " NULL", " DEFAULT ", " AUTO_INCREMENT", " PRIMARY KEY", " UNIQUE"]
            split_pos = len(remainder)
            for keyword in keywords:
                idx = remainder.find(keyword)
                if idx != -1:
                    split_pos = min(split_pos, idx)
            data_type = remainder[:split_pos].strip()
            nullable = "否" if "NOT NULL" in remainder else "是"
            default = "-"
            m_default = re.search(r"DEFAULT\s+(.+?)(?:\s+ON UPDATE|\s*$)", remainder)
            if m_default:
                default = m_default.group(1).strip().strip(",")
            if "AUTO_INCREMENT" in remainder:
                constraint_tokens.append("自增")
            if "PRIMARY KEY" in remainder:
                constraint_tokens.append("主键")
            columns.append(Column(name=name, data_type=data_type, nullable=nullable, default=default, constraints=constraint_tokens))
        for constraint in table_constraints:
            if constraint.startswith("PRIMARY KEY"):
                cols = re.findall(r"\((.*?)\)", constraint)
                if cols:
                    names = [x.strip().strip("`") for x in cols[0].split(",")]
                    for col in columns:
                        if col.name in names and "主键" not in col.constraints:
                            col.constraints.append("主键")
            elif constraint.startswith("UNIQUE KEY"):
                cols = re.findall(r"\((.*?)\)", constraint)
                if cols:
                    names = [x.strip().strip("`") for x in cols[0].split(",")]
                    for col in columns:
                        if col.name in names and "唯一" not in col.constraints:
                            col.constraints.append("唯一")
            elif constraint.startswith("CONSTRAINT") and "FOREIGN KEY" in constraint:
                m_fk = re.search(r"FOREIGN KEY \((.*?)\) REFERENCES (\w+) \((.*?)\)", constraint)
                if m_fk:
                    cols, ref_table, ref_cols = m_fk.groups()
                    col_names = [x.strip().strip("`") for x in cols.split(",")]
                    for col in columns:
                        if col.name in col_names:
                            col.constraints.append(f"外键->{ref_table}({ref_cols})")
        tables.append(TableDef(table_name, columns, table_constraints))
    return tables


def field_desc(table_name: str, column_name: str) -> str:
    specific_key = f"{table_name}.{column_name}"
    if specific_key in FIELD_DESC_OVERRIDE:
        return FIELD_DESC_OVERRIDE[specific_key]
    if column_name in GENERIC_FIELD_DESC:
        return GENERIC_FIELD_DESC[column_name]
    return "业务字段，结合所在实体承载对应场景数据。"


def update_cover(doc: Document) -> None:
    replacements = {
        0: "101-概要设计说明书",
        4: "校园活动智能推荐与组队平台",
        5: "概要设计说明书",
        22: "二〇二六年四月",
        28: "核    定：指导教师",
        29: "审    查：项目负责人",
        30: "校    核：系统设计负责人",
        31: "编    制：第25组项目组",
        45: "文档编辑记录",
    }
    for idx, text in replacements.items():
        if idx < len(doc.paragraphs):
            doc.paragraphs[idx].text = text
            for run in doc.paragraphs[idx].runs:
                size = 12
                bold = False
                font = "宋体"
                if idx == 4:
                    size = 20
                    bold = True
                    font = "黑体"
                elif idx == 5:
                    size = 24
                    bold = True
                    font = "黑体"
                elif idx in {28, 29, 30, 31}:
                    size = 14
                elif idx == 22:
                    size = 14
                elif idx == 0:
                    size = 14
                    bold = True
                    font = "黑体"
                elif idx == 45:
                    size = 16
                    bold = True
                    font = "黑体"
                set_run_font(run, size, bold, east_asia=font)
    table = doc.tables[0]
    set_table_borders(table)
    row = table.rows[1]
    values = ["1", "2026-04-20", "第25组项目组", "依据需求规格说明书、系统原型与后端实现编制概要设计说明书", "V1.0", "100%"]
    for cell, value in zip(row.cells, values):
        cell.text = value
        style_cell(cell, center=False)
    if len(table.rows) > 2:
        tbl = table._tbl
        tbl.remove(table.rows[2]._tr)


def render_diagrams() -> dict[str, Path]:
    PUML_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    out_paths: dict[str, Path] = {}
    for name, content in DIAGRAMS.items():
        puml_path = PUML_DIR / name
        puml_path.write_text(content, encoding="utf-8")
        subprocess.run([
            "java",
            "-Djava.awt.headless=true",
            "-jar",
            str(PLANTUML_JAR),
            "-charset",
            "UTF-8",
            "-tpng",
            "-o",
            str(ASSET_DIR),
            str(puml_path),
        ], check=True, cwd=ROOT)
        out_paths[name] = ASSET_DIR / name.replace(".puml", ".png")
    return out_paths


def add_picture(doc: Document, image_path: Path, caption: str, width_cm: float = 15.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_caption(doc, caption)


def build_document(diagrams: dict[str, Path], tables: list[TableDef]) -> None:
    doc = Document(str(TEMPLATE))
    set_update_fields(doc)
    update_cover(doc)
    clear_after(doc, 46)

    doc.add_page_break()
    add_heading(doc, "目  录", 1)
    p = doc.add_paragraph()
    add_toc_field(p)
    apply_para_format(p)
    p.paragraph_format.first_line_indent = Cm(0)

    doc.add_page_break()

    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 目的", 2)
    add_body_paragraph(doc, "本说明书用于描述“校园活动智能推荐与组队平台”的概要设计方案，明确系统总体结构、功能划分、接口组织、数据库设计、安全与维护策略，为后续详细设计、编码实现、测试验证和课程答辩提供统一依据。本文档的主要读者包括课程指导教师、项目组成员、前后端开发人员、测试人员及系统维护人员。")
    add_heading(doc, "1.2 背景", 2)
    add_body_paragraph(doc, "校园活动智能推荐与组队平台面向高校在校学生、活动组织者和系统管理员，旨在解决校园活动信息分散、兴趣匹配不足、组队效率低和协同沟通不便等问题。系统采用 Web 用户平台与 Web 管理后台相结合的方式，对外提供活动浏览、智能推荐、活动报名、组队协作、即时消息、通知提醒、智能客服问答和后台审核运营等能力；对内通过 Spring Boot 后端、MySQL 数据库、邮件服务、AI 外部服务与 Dify 知识库 Chatflow 协同运行。项目任务提出者为课程教学团队，开发者为第25组项目组，最终用户包括普通学生、活动组织者与平台管理员。")
    add_heading(doc, "1.3 术语和定义", 2)
    add_table(doc, ["术语", "说明"], [
        ("兴趣画像", "基于用户显式兴趣标签、点击收藏行为、历史参加活动记录和历史组队参与信息构建的用户偏好表达。"),
        ("智能推荐", "结合活动标签、用户兴趣、历史参加活动、历史组队参与、运营权重和时间因素对活动进行综合排序并生成推荐结果。"),
        ("智能组队", "基于活动上下文、用户兴趣与队伍状态，为用户提供创建队伍、申请入队和候选协作者推荐的能力。"),
        ("活动审核", "管理员对活动内容、状态和展示资格进行审核与管控的业务流程。"),
        ("SSE 实时通信", "Server-Sent Events，一种由服务端向浏览器持续推送事件消息的实时通信机制。"),
        ("RBAC 权限控制", "Role-Based Access Control，基于角色的访问控制模型，用于区分普通用户、组织者和管理员的操作权限。"),
        ("AI 标签提取", "通过外部大模型服务对活动标题和描述进行语义分析，自动提取活动标签或推荐辅助信息。"),
        ("Dify Chatflow", "基于知识库构建的智能客服流程，用于解答用户在平台使用过程中的常见问题和操作疑问。"),
    ], widths=[4, 12])
    add_heading(doc, "1.4 参考资料", 2)
    refs = [
        "《校园活动智能推荐与组队平台需求规格说明书》，第25组项目组，2026年3月。",
        "《101-概要设计说明书(模版)》，课程提供模板文件。",
        "Spring Boot Reference Documentation 3.1.4，VMware, Inc.",
        "MySQL 8.0 Reference Manual，Oracle Corporation。",
        "PlantUML Language Reference Guide，用于架构图与数据关系图绘制。",
        "SMTP 邮件服务接口文档（项目中使用 QQ 邮箱 SMTP 服务配置）。",
        "AI 外部服务接口说明，用于活动标签提取、推荐辅助和智能组队建议能力设计。",
        "Dify Chatflow 接口与知识库配置说明，用于智能客服问答能力设计。"
    ]
    for idx, ref in enumerate(refs, 1):
        add_body_paragraph(doc, f"{idx}. {ref}")

    add_heading(doc, "2 系统概述", 1)
    add_heading(doc, "2.1 需求规定", 2)
    add_inline_title(doc, "功能要求：", "系统应支持用户注册登录、兴趣标签维护、活动浏览检索、活动发布与编辑、收藏与报名、队伍创建与申请、队伍聊天、活动群聊、用户私聊、通知中心、个人中心、智能客服以及后台审核与运营管理。")
    add_inline_title(doc, "性能要求：", "活动列表、推荐流和通知列表应在常规校园网络环境下快速响应；推荐接口响应时间原则上控制在 2 秒以内，普通分页查询应在 1 秒级返回。")
    add_inline_title(doc, "输入输出要求：", "输入包括注册信息、活动内容、标签、报名信息、聊天消息、上传文件、后台审核操作以及智能客服问答请求；输出包括统一 JSON 响应数据、推荐列表、消息推送事件、通知、客服回答和管理统计结果。")
    add_inline_title(doc, "数据管理能力要求：", "系统需持久化用户、活动、标签、报名、组队、消息、通知、行为日志和运营推荐位等核心数据，并基于用户标签、历史参加活动与组队信息构建推荐输入。")
    add_inline_title(doc, "故障处理能力要求：", "系统在 AI 服务、Dify 客服服务或邮件服务不可用时应具备降级能力，在接口异常、状态冲突或资源不存在时向前端返回清晰错误信息。")
    add_heading(doc, "2.2 开发及运行环境", 2)
    add_table(doc, ["类别", "配置说明"], [
        ("开发语言", "Java 17、HTML5、CSS3、JavaScript"),
        ("后端框架", "Spring Boot 3.1.4、Spring JDBC、Spring Mail"),
        ("数据库", "MySQL 8.x"),
        ("构建工具", "Maven"),
        ("运行方式", "Spring Boot Jar 独立部署，浏览器访问 Web 页面"),
        ("前端运行环境", "Chrome、Edge、Safari 等现代浏览器"),
        ("消息机制", "HTTP + REST/JSON；SSE 实时推送"),
        ("第三方能力", "SMTP 邮件服务、AI 外部服务、Dify 知识库 Chatflow"),
        ("文件存储", "项目静态目录下的本地文件存储"),
    ], widths=[4, 12])
    add_heading(doc, "2.3 限制和约束", 2)
    constraints = [
        "系统业务场景基于校园用户身份，账号、学号、学院等信息应符合校园活动管理实际。",
        "推荐结果依赖兴趣标签、历史参加活动、历史组队参与信息、行为日志和标签体系的完整性；冷启动用户默认采用通用热门和时间优先策略。",
        "AI 能力和 Dify 智能客服作为正式外部服务纳入设计，但应提供默认标签、规则排序、常见问题静态答复和人工处理兜底机制。",
        "实时消息依赖浏览器 SSE 长连接和稳定网络环境，在弱网场景下允许前端回退到主动刷新。",
        "文件上传采用本地静态目录存储，需限制文件类型并规范访问路径。",
        "涉及活动审核、报名审批和队伍审批的接口必须经过登录校验和资源权限校验。",
    ]
    for item in constraints:
        add_body_paragraph(doc, f"（1）{item}" if item == constraints[0] else f"（{constraints.index(item)+1}）{item}")

    add_heading(doc, "3 总体设计", 1)
    add_heading(doc, "3.1 设计原则", 2)
    principles = [
        "模块化原则：按照认证、活动、推荐、组队、消息、通知、后台和支撑能力拆分模块，降低耦合度。",
        "统一性原则：全部 API 使用统一响应模型 ApiResponse<T>，统一异常处理策略和权限控制入口。",
        "安全优先原则：以 Token 鉴权、角色权限、输入校验和敏感信息保护作为基础设计要求。",
        "可扩展原则：推荐算法、AI 能力、上传存储和运营策略均采用可替换、可扩展的设计口径。",
        "可维护原则：通过配置分离、定时清理、统一日志与模块化控制器组织提高可维护性。",
    ]
    for idx, item in enumerate(principles, 1):
        add_body_paragraph(doc, f"{idx}. {item}")
    add_heading(doc, "3.2 总体架构", 2)
    add_body_paragraph(doc, "系统总体上划分为前端展示层、业务服务层、数据持久层和外部服务层四个部分。Web 用户平台和 Web 管理后台通过 REST/JSON 与后端交互，在聊天场景下通过 SSE 订阅实时消息；后端通过 Spring Boot 统一承载认证、活动、推荐、组队、消息、通知、智能客服和后台运营等业务；MySQL 负责核心数据存储；文件存储、SMTP 邮件服务、AI 外部服务和 Dify 知识库 Chatflow 为系统提供支撑能力。系统总体架构如图3-1所示。")
    add_picture(doc, diagrams["overall-architecture.puml"], "图3-1 系统总体架构图")
    add_body_paragraph(doc, "总体架构中，前端层主要负责界面展示和用户交互，业务服务层负责权限校验、流程编排和数据计算，数据层负责持久化与查询，外部服务层负责验证码投递、智能标签提取、推荐辅助和知识库客服问答等能力。该结构能够同时满足课程项目的实现复杂度与后续功能扩展需求。")
    add_heading(doc, "3.3 功能架构", 2)
    add_body_paragraph(doc, "系统功能架构围绕“发现活动、参与活动、发起组队、协同沟通、后台治理”主线展开。功能架构如图3-2所示，功能需求与模块分配矩阵见表3-1。")
    add_picture(doc, diagrams["functional-architecture.puml"], "图3-2 功能架构图")
    add_caption(doc, "表3-1 功能需求-模块分配矩阵")
    add_table(doc, ["功能需求", "用户与认证", "活动管理", "智能推荐", "组队协作", "消息通信", "通知与个人中心", "管理后台", "平台支撑"], MODULE_MATRIX_ROWS, widths=[4.2, 1.4, 1.4, 1.4, 1.4, 1.4, 1.8, 1.4, 1.4])

    module_content = [
        ("3.3.1 用户与认证模块", [
            ("3.3.1.1 账号注册登录功能：", "提供账号密码注册登录、邮箱验证码注册登录、邮箱绑定与修改密码等能力，核心类包括 AuthController、TokenService、PasswordHasher、EmailCodeService、SmsCodeService。"),
            ("3.3.1.2 兴趣画像维护功能：", "通过标签接口与用户兴趣维护接口管理用户显式兴趣标签，并结合历史参加活动、收藏点击和组队参与信息形成推荐排序输入。"),
        ]),
        ("3.3.2 活动发现与活动管理模块", [
            ("3.3.2.1 活动浏览与详情功能：", "支持活动列表分页、关键词检索、标签筛选、活动高亮和活动详情展示，核心控制器为 ActivityController。"),
            ("3.3.2.2 活动发布与报名审批功能：", "支持组织者发布活动、编辑活动、查看报名记录、审批报名以及获取组织者联系方式，并通过 notifications 向相关用户推送状态变更消息。"),
        ]),
        ("3.3.3 智能推荐模块", [
            ("3.3.3.1 首页推荐流功能：", "RecommendController 基于活动标签、用户兴趣、历史参加活动、历史组队参与、行为日志、时间窗口与运营权重生成推荐结果。"),
            ("3.3.3.2 AI 推荐增强功能：", "通过外部 AI 服务对活动内容进行语义分析和标签补充，在外部服务失效时退化到规则排序与现有标签体系。"),
        ]),
        ("3.3.4 组队协作模块", [
            ("3.3.4.1 队伍管理功能：", "TeamController 负责队伍列表、队伍详情、创建、修改、删除与成员展示。"),
            ("3.3.4.2 入队申请审批功能：", "系统保存入队申请记录，队长可审批通过或拒绝，并在成功后维护 team_member 与通知数据。"),
        ]),
        ("3.3.5 消息通信模块", [
            ("3.3.5.1 聊天消息功能：", "ChatController、DirectMessageController、ActivityChatController 分别支撑队伍聊天、用户私聊和活动群聊。"),
            ("3.3.5.2 实时推送功能：", "ChatRealtimeController 和 ChatRealtimeService 基于 SSE 为三类聊天场景提供实时消息推送和订阅权限校验。"),
        ]),
        ("3.3.6 通知与个人中心模块", [
            ("3.3.6.1 通知中心功能：", "NotificationController 支持通知分页查询、单条已读、全部已读和删除，同时定时清理过期通知。"),
            ("3.3.6.2 个人中心与智能客服功能：", "ProfileController 提供资料维护、头像设置、我的收藏、我的报名、我的队伍和聊天入口聚合能力，同时系统在用户侧提供基于 Dify 知识库 Chatflow 的智能客服问答入口，用于解答平台使用问题。"),
        ]),
        ("3.3.7 管理后台模块", [
            ("3.3.7.1 审核治理功能：", "AdminController 负责用户状态管理、活动审核、活动状态调整和统计看板。"),
            ("3.3.7.2 运营配置功能：", "后台可以维护标签体系和首页精选活动推荐位，直接影响前台展示与推荐排序。"),
        ]),
        ("3.3.8 平台支撑模块", [
            ("3.3.8.1 文件上传功能：", "UploadController 支持头像与封面上传，并统一校验文件类型、目录生成和访问路径返回。"),
            ("3.3.8.2 公共支撑功能：", "Api 全局响应与异常处理、ApiAuthConfig 鉴权拦截器、Bootstrap 初始化脚本和 ExpiryCleanupJob 定时清理任务共同构成平台支撑能力。"),
        ]),
    ]
    for title, items in module_content:
        add_heading(doc, title, 3)
        for inline_title, body in items:
            add_inline_title(doc, inline_title, body)

    add_heading(doc, "3.4 技术架构", 2)
    add_table(doc, ["层次", "主要技术", "职责说明"], [
        ("前端展示层", "HTML5、CSS3、JavaScript", "负责页面渲染、表单交互、消息订阅和后台管理界面展示。"),
        ("接口控制层", "Spring Web Controller", "负责路由映射、参数解析、权限校验入口和统一响应封装。"),
        ("业务编排层", "控制器内业务编排 + 公共服务类", "负责登录认证、推荐排序、聊天推送、邮件验证码、智能客服转发和审核通知等流程处理。"),
        ("数据访问层", "JdbcTemplate + MySQL", "负责核心业务数据的增删改查、统计查询和事务控制。"),
        ("实时通信层", "SSE（Server-Sent Events）", "负责聊天消息的实时推送和频道级权限控制。"),
        ("安全支撑层", "TokenService、PasswordHasher、ApiAuthConfig", "负责密码哈希、Token 签发解析和 RBAC 访问控制。"),
        ("外部集成层", "SMTP、AI HTTP 接口、Dify Chatflow、本地文件 I/O", "负责邮件验证码、AI 能力调用、知识库客服问答和文件上传持久化。"),
    ], widths=[3.2, 5.1, 7.2])
    add_heading(doc, "3.5 数据架构", 2)
    add_body_paragraph(doc, "系统数据架构按照业务域划分为用户域、活动域、组队域、消息域、运营域和行为域六类。用户域保存账号信息、兴趣画像和联系方式；活动域保存活动主数据、活动标签和报名记录；组队域保存队伍、成员与申请；消息域保存队伍消息、活动群聊、私聊及通知；运营域保存精选推荐位和后台标签；行为域保存点击、收藏、报名等事件日志。推荐模块重点使用 user_interest、event_log、registrations、team_member 与 teams 等数据构建画像与排序输入。核心数据实体关系如图3-3所示。")
    add_picture(doc, diagrams["data-architecture.puml"], "图3-3 核心数据实体关系图")
    add_body_paragraph(doc, "在数据组织上，系统以 users、activities、teams 为核心主实体，通过多对多关联表和日志表构建推荐、沟通和运营链路。主外键约束保证了报名、收藏、组队、消息和通知等数据在删除或更新场景下的一致性；Dify 智能客服知识库作为外部知识资产，不落地到本系统业务主表中。")
    add_heading(doc, "3.6 部署架构", 2)
    add_body_paragraph(doc, "系统部署采用轻量化单体后端加静态前端的结构。浏览器直接访问前端页面并调用 API 服务；后端服务以 Spring Boot Jar 形式运行，连接 MySQL 数据库和本地文件目录，同时通过网络访问 SMTP 邮件服务、AI 外部服务和 Dify Chatflow 服务。部署架构如图3-4所示。")
    add_picture(doc, diagrams["deployment-architecture.puml"], "图3-4 部署架构图")
    add_body_paragraph(doc, "该部署方式适合课程项目和校园内部试运行环境，具备配置简单、部署成本低和便于演示的特点；若后续系统规模扩大，可进一步演进为前后端分离部署和云对象存储模式。")

    add_heading(doc, "4 数据结构设计", 1)
    add_body_paragraph(doc, "本章依据数据库脚本 schema.sql 对系统核心数据结构进行说明。为保证数据库设计完整呈现，先给出全部业务表清单，再逐表给出字段、数据类型及关键约束说明。")
    add_caption(doc, "表4-1 数据库表清单")
    add_table(doc, ["序号", "表名", "中文名称", "主要用途"], [
        (str(i), table_def.name, TABLE_CN.get(table_def.name, table_def.name), TABLE_PURPOSE.get(table_def.name, "业务数据表"))
        for i, table_def in enumerate(tables, 1)
    ], widths=[1.2, 3.2, 3.2, 8.6], body_font_size=9.5, header_font_size=9.5)
    table_no = 1
    for table_def in tables:
        if table_def.name not in TABLE_CN:
            continue
        add_inline_title(doc, f"表4-{table_no + 1} {table_def.name}（{TABLE_CN[table_def.name]}）：", TABLE_PURPOSE[table_def.name])
        rows = []
        for column in table_def.columns:
            constraints = "、".join(column.constraints) if column.constraints else "-"
            rows.append([
                column.name,
                column.data_type,
                "是" if "主键" in column.constraints else "否",
                column.nullable,
                column.default,
                constraints,
                field_desc(table_def.name, column.name),
            ])
        add_table(
            doc,
            ["字段名", "数据类型", "主键", "可空", "默认值", "关键约束", "字段说明"],
            rows,
            widths=[2.2, 2.3, 1.0, 1.0, 1.6, 2.6, 5.6],
            body_font_size=8.2,
            header_font_size=8.8,
        )
        table_no += 1

    add_heading(doc, "5 接口设计", 1)
    add_body_paragraph(doc, "系统接口采用 RESTful JSON 与 SSE 相结合的方式。除文件上传接口外，HTTP 接口默认以 JSON 作为报文体；统一响应模型定义为 ApiResponse<T> = { success, message, data }。鉴权接口之外的受保护接口通过 Authorization: Bearer <token> 传递登录凭证。")
    add_heading(doc, "5.1 内部接口", 2)
    for idx, (name, module, items) in enumerate(INTERNAL_INTERFACES, 1):
        add_inline_title(doc, f"（{idx}）{name}：", f"该类接口主要服务于{module}。")
        add_table(doc, ["方法", "接口路径", "说明"], items, widths=[2.2, 6.5, 8.3])
    add_heading(doc, "5.2 外部接口", 2)
    add_table(doc, ["外部接口", "协议/方式", "接口说明", "主要用途"], EXTERNAL_INTERFACES, widths=[4.2, 3.3, 5.8, 4.2])
    add_body_paragraph(doc, "在外部接口设计中，AI 服务与 Dify Chatflow 均为正式设计能力。系统可将活动标题、活动描述、用户画像及活动上下文发送至 AI 服务，获取标签提取、推荐辅助或组队建议结果；同时可将用户问题、页面上下文和知识库检索条件发送至 Dify Chatflow，获取智能客服回答。若外部服务异常，则分别退化为使用已有标签、规则排序、常见问题静态答复或人工补录方式。")

    add_heading(doc, "6 出错处理设计", 1)
    add_heading(doc, "6.1 出错信息", 2)
    add_table(doc, ["错误类别", "典型场景", "输出形式", "处理方法"], ERROR_ROWS, widths=[3, 4.5, 5.2, 5.3])
    add_heading(doc, "6.2 出错处理对策", 2)
    strategies = [
        "统一异常处理：通过 Api.GlobalExceptionHandler 对业务异常、参数异常和未知异常进行统一封装，减少前端处理分支。",
        "状态冲突保护：对于重复报名、重复申请、审核状态不合法等场景，通过数据库唯一约束和业务判断联合防护。",
        "AI 服务降级：AI 标签提取、推荐增强或组队建议失败时，系统退化为本地规则、默认标签和既有推荐逻辑。",
        "Dify 客服降级：知识库 Chatflow 不可用时，系统退化为常见问题静态答复、人工联系方式提示或留言收集机制。",
        "消息补偿：实时推送失败不影响消息入库，用户可在重新进入聊天页面时通过历史接口补齐消息。",
        "通知清理：通知中心通过查询前清理和定时任务清理相结合的方式删除过期通知，避免数据无界增长。",
        "备份与恢复：数据库采用常规备份策略，上传文件目录需定期备份，以便在误删或环境迁移时恢复。",
    ]
    for idx, item in enumerate(strategies, 1):
        add_body_paragraph(doc, f"{idx}. {item}")

    add_heading(doc, "7 安全保密设计", 1)
    security_items = [
        "身份认证安全：登录后由 TokenService 签发访问令牌，受保护接口通过 ApiAuthConfig 统一鉴权。",
        "口令与验证码安全：密码以单向哈希形式保存，邮箱/短信验证码使用摘要值存储并设置有效期、冷却时间和尝试次数。",
        "角色与权限控制：系统按普通用户、组织者、管理员进行 RBAC 控制，确保活动审核、用户管理和运营配置只能由管理员执行。",
        "输入校验与内容安全：对账号、邮箱、上传文件、报名信息和聊天内容进行基础校验，降低注入、越权和无效数据风险。",
        "敏感信息保护：手机号、邮箱等联系方式仅在必要场景下返回，组织者联系方式查询受业务流程控制。",
        "数据完整性保护：数据库通过主键、唯一键和外键约束维护实体一致性，避免悬挂数据和重复数据。",
        "日志与审计：通过行为日志、通知记录、审核记录和系统异常日志支撑问题追踪和运维审计。",
    ]
    for idx, item in enumerate(security_items, 1):
        add_body_paragraph(doc, f"{idx}. {item}")

    add_heading(doc, "8 维护设计", 1)
    maintenance_items = [
        "配置分离：数据库、邮件、Token、验证码等运行参数集中存放在 application.yml 中，便于按环境调整。",
        "初始化与演示数据：Bootstrap 和 SQL 初始化脚本负责建表、补字段与演示数据加载，便于课程演示和快速部署。",
        "模块化组织：系统按 activity、auth、team、message、notification、profile、admin 等包组织代码，降低维护复杂度。",
        "统一响应与异常处理：公共 Api 类减少接口风格不一致问题，提升前后端联调效率。",
        "定时维护任务：ExpiryCleanupJob 自动清理过期通知，避免历史数据对查询性能的持续影响。",
        "可扩展算法与 AI 适配：推荐逻辑、标签提取、组队建议和 Dify 客服接入均可在保持接口不变的前提下替换内部实现，支持后续版本演进。",
    ]
    for idx, item in enumerate(maintenance_items, 1):
        add_body_paragraph(doc, f"{idx}. {item}")

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    doc.save(str(OUTPUT))


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"模板不存在: {TEMPLATE}")
    if not SCHEMA.exists():
        raise FileNotFoundError(f"数据库脚本不存在: {SCHEMA}")
    if not PLANTUML_JAR.exists():
        raise FileNotFoundError(f"PlantUML 运行文件不存在: {PLANTUML_JAR}")
    diagrams = render_diagrams()
    tables = parse_schema(SCHEMA.read_text(encoding="utf-8"))
    build_document(diagrams, tables)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
